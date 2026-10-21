"""用 LibreOffice 重算 xlsx 里的全部公式，原地写回，并以 JSON 报告错误值。

    python3 recalc.py 文件.xlsx [超时秒数]      默认 60 秒
    python3 recalc.py 文件.xlsx --force         有外部链接也照算（接受链接丢失）

openpyxl 写进去的公式没有缓存值，pandas、data_only=True 和大多数预览读出来
全是空。重算之后才有值，也才知道有没有 #REF!、#DIV/0! 这类错误。

输出 JSON：
    {"status": "success" | "errors_found", "total_formulas": N,
     "total_errors": N, "error_summary": {"#REF!": {"count": N, "locations": [...]}}}
只有 JSON 里是 "error"（没能重算）时退出码才非 0；errors_found 退出码是 0，
所以**不要拿退出码当作表格没问题**，要看 status。
"""

import json
import os
import shutil
import subprocess
import sys
import tempfile
import zipfile

ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!")
MAX_LOCATIONS = 100


def fail(message):
    print(json.dumps({"error": message}, ensure_ascii=False))
    return 1


def has_external_links(path):
    with zipfile.ZipFile(path) as archive:
        return any(name.startswith("xl/externalLinks/") for name in archive.namelist())


def main(argv=None):
    args = list(sys.argv[1:] if argv is None else argv)
    force = "--force" in args
    args = [a for a in args if a != "--force"]
    if not args:
        print(__doc__)
        return 2
    path = args[0]
    timeout = int(args[1]) if len(args) > 1 else 60
    if not os.path.exists(path):
        return fail(f"文件不存在：{path}")
    if has_external_links(path) and not force:
        return fail(
            "这个工作簿引用了别的文件（外部链接）。那些文件不在这里，重算会把相关格子变成 "
            "#NAME? 并删掉链接。先把这些格子的值从原文件抄出来；确认可以丢掉链接再加 --force。"
        )

    here = os.path.dirname(os.path.abspath(__file__))
    out_dir = tempfile.mkdtemp(prefix="recalc-")
    try:
        try:
            subprocess.run(
                [sys.executable, os.path.join(here, "soffice.py"), "--headless",
                 "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", out_dir, path],
                capture_output=True, text=True, timeout=timeout,
            )
        except subprocess.TimeoutExpired:
            return fail(f"LibreOffice {timeout} 秒内没有算完")
        produced = os.path.join(out_dir, os.path.basename(path))
        if not os.path.exists(produced):
            return fail("LibreOffice 没能打开这个工作簿")

        from openpyxl import load_workbook

        formulas = load_workbook(produced)
        values = load_workbook(produced, data_only=True)
        total_formulas = 0
        summary = {}
        for sheet in formulas.worksheets:
            computed = values[sheet.title]
            for row in sheet.iter_rows():
                for cell in row:
                    if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                        continue
                    total_formulas += 1
                    value = computed[cell.coordinate].value
                    if isinstance(value, str) and value in ERRORS:
                        entry = summary.setdefault(value, {"count": 0, "locations": []})
                        entry["count"] += 1
                        if len(entry["locations"]) < MAX_LOCATIONS:
                            entry["locations"].append(f"{sheet.title}!{cell.coordinate}")
        shutil.copyfile(produced, path)
    finally:
        shutil.rmtree(out_dir, ignore_errors=True)

    total_errors = sum(entry["count"] for entry in summary.values())
    for entry in summary.values():
        withheld = entry["count"] - len(entry["locations"])
        if withheld:
            entry["locations_truncated"] = withheld
    print(json.dumps({
        "status": "errors_found" if total_errors else "success",
        "total_formulas": total_formulas,
        "total_errors": total_errors,
        "error_summary": summary,
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
