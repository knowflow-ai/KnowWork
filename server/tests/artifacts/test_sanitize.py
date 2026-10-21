"""表格公式注入防护。"""

import io
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from knowwork_server.artifacts.sanitize import SpreadsheetRejected, escape_cell, sanitize_csv, sanitize_xlsx

from ..office_samples import xlsx


@pytest.mark.parametrize(("raw", "escaped"), [
    ("=1+1", "'=1+1"), ("+cmd|' /C calc'!A0", "'+cmd|' /C calc'!A0"), ("-2+3", "'-2+3"), ("@SUM(A1)", "'@SUM(A1)"),
    ("\t=x", "'\t=x"), ("\r=x", "'\r=x"), ("-5", "-5"), ("+3.2", "+3.2"), ("-1e5", "-1e5"), ("普通", "普通"),
    ("", ""), ("a=b", "a=b"),
])
def test_escape_cell(raw, escaped):
    assert escape_cell(raw) == escaped


def test_csv_escapes_and_keeps_encoding():
    data = "名称,值\n=HYPERLINK(\"http://x\"),-5\n@a,ok\n".encode("gb18030")
    out, changed = sanitize_csv(data)
    assert changed
    text = out.decode("gb18030")
    assert "'=HYPERLINK" in text and "'@a" in text and ",-5" in text and "名称" in text


def test_csv_untouched_and_bom_kept():
    clean = b"a,b\n1,2\n"
    assert sanitize_csv(clean) == (clean, False)
    out, changed = sanitize_csv("﻿a\t=1\n".encode(), delimiter="\t")
    assert changed and out.startswith(b"\xef\xbb\xbf") and b"\t'=1" in out


def test_csv_rejects_unknown_encoding_and_bad_csv():
    with pytest.raises(SpreadsheetRejected):
        sanitize_csv(b"\xff\xfe\xfd\x81\x30")
    with pytest.raises(SpreadsheetRejected):
        sanitize_csv(b'"' + b"x" * 200_000 + b'"\n')  # 超过 csv 模块的单字段上限


def test_xlsx_formulas_become_quoted_text():
    data = xlsx([["标题", "=SUM(1,2)"], ["@cmd", -5], ["+x", "普通"]])
    out, changed = sanitize_xlsx(data)
    assert changed
    sheet = load_workbook(io.BytesIO(out)).active
    assert (sheet["B1"].value, sheet["B1"].data_type, sheet["B1"].quotePrefix) == ("=SUM(1,2)", "s", True)
    assert sheet["A2"].quotePrefix and sheet["A3"].quotePrefix
    assert sheet["B2"].value == -5 and not sheet["B2"].quotePrefix
    assert sheet["A1"].value == "标题" and not sheet["A1"].quotePrefix


def test_xlsx_without_risky_cells_is_returned_byte_for_byte():
    data = xlsx([["a", 1], ["b", 2]])
    assert sanitize_xlsx(data) == (data, False)


def test_xlsx_array_formula():
    from openpyxl.worksheet.formula import ArrayFormula

    workbook = Workbook()
    workbook.active["A1"] = ArrayFormula("A1:A1", "=SUM(B1:B2)")
    buffer = io.BytesIO()
    workbook.save(buffer)
    out, changed = sanitize_xlsx(buffer.getvalue())
    assert changed and load_workbook(io.BytesIO(out)).worksheets[0]["A1"].value == "=SUM(B1:B2)"


def test_xlsx_rejects_broken_and_bombs(monkeypatch):
    with pytest.raises(SpreadsheetRejected):
        sanitize_xlsx(b"PK\x03\x04garbage")
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("x.xml", "<x/>")
    with pytest.raises(SpreadsheetRejected):  # 是 zip 但不是工作簿
        sanitize_xlsx(buffer.getvalue())
    monkeypatch.setattr("knowwork_server.artifacts.sanitize.MAX_XLSX_UNCOMPRESSED", 10)
    with pytest.raises(SpreadsheetRejected, match="过大"):
        sanitize_xlsx(xlsx([["a"]]))
